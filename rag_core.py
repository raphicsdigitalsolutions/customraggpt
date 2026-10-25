"""Document ingestion and citation helpers for NEXORA's RAG workflow."""

import csv
import io
import json
import os
import tempfile
from pathlib import Path
from typing import Any, Callable

from langchain_community.document_loaders import PyPDFLoader
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_openai import OpenAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter


def load_file(file_name: str, content: bytes) -> list[Document]:
    """Read a supported file and return LangChain Documents."""
    suffix = Path(file_name).suffix.lower()

    if suffix == ".pdf":
        tmp_path = None
        try:
            with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                tmp.write(content)
                tmp_path = tmp.name
            docs = PyPDFLoader(tmp_path).load()
            for doc in docs:
                doc.metadata["source"] = file_name
            return docs
        finally:
            if tmp_path and os.path.exists(tmp_path):
                os.unlink(tmp_path)

    if suffix in {".txt", ".md"}:
        text = content.decode("utf-8-sig", errors="replace")
        return [Document(page_content=text, metadata={"source": file_name, "page": 0})] if text.strip() else []

    if suffix == ".csv":
        text = content.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
        if not rows:
            return []
        header = " | ".join(rows[0])
        docs = []
        for row_num, row in enumerate(rows[1:], start=2):
            row_text = " | ".join(row)
            if row_text.strip():
                docs.append(
                    Document(
                        page_content=f"CSV columns: {header}\nRow {row_num}: {row_text}",
                        metadata={"source": file_name, "page": row_num - 1, "row": row_num},
                    )
                )
        if not docs:
            docs = [Document(page_content=f"CSV columns: {header}", metadata={"source": file_name, "page": 0})]
        return docs

    if suffix == ".json":
        try:
            parsed = json.loads(content.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValueError(f"{file_name} is not valid JSON: {exc}") from exc

        pretty_json = json.dumps(parsed, ensure_ascii=False, indent=2)
        if not pretty_json.strip():
            return []
        return [
            Document(
                page_content=f"JSON file: {file_name}\n\n{pretty_json}",
                metadata={"source": file_name, "page": 0},
            )
        ]

    if suffix == ".docx":
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise RuntimeError(
                "DOCX support requires python-docx. Install it with: pip install python-docx"
            ) from exc

        docx = DocxDocument(io.BytesIO(content))
        parts = [p.text.strip() for p in docx.paragraphs if p.text.strip()]
        for table in docx.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    parts.append(" | ".join(values))
        full_text = "\n".join(parts)
        return [Document(page_content=full_text, metadata={"source": file_name, "page": 0})] if full_text.strip() else []

    raise ValueError(f"Unsupported file type: {suffix}")


def build_vectorstore(
    files: list[Any],
    size: int,
    overlap: int,
    api_key: str,
    on_warning: Callable[[str], None] | None = None,
):
    """Parse and index multiple uploaded files, retaining their source metadata."""
    if overlap >= size:
        raise ValueError("Chunk overlap must be smaller than chunk size.")

    all_docs: list[Document] = []
    file_names: list[str] = []
    total_pages = 0

    for uploaded in files:
        docs = load_file(uploaded.name, uploaded.getvalue())
        docs = [doc for doc in docs if doc.page_content and doc.page_content.strip()]
        if not docs:
            if on_warning:
                on_warning(f"No readable text found in {uploaded.name}; skipped.")
            continue
        file_names.append(uploaded.name)
        for doc in docs:
            doc.metadata["source"] = uploaded.name
        all_docs.extend(docs)
        total_pages += len(docs)

    if not all_docs:
        raise ValueError(
            "No readable text was found in the selected files. Scanned PDFs may need OCR first."
        )

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=size,
        chunk_overlap=overlap,
        separators=["\n\n", "\n", " ", ""],
    )
    chunks = splitter.split_documents(all_docs)
    if not chunks:
        raise ValueError("No searchable chunks could be created.")

    embeddings = OpenAIEmbeddings(model="text-embedding-3-small", api_key=api_key)
    return FAISS.from_documents(chunks, embeddings), len(chunks), total_pages, file_names


def display_page(metadata: dict[str, Any]) -> str:
    page = metadata.get("page")
    if isinstance(page, int):
        if metadata.get("row") is not None:
            return f"row {metadata['row']}"
        return str(page + 1)
    return str(page) if page is not None else "—"


def format_docs(docs: list[Document]) -> str:
    formatted = []
    for doc in docs:
        source = doc.metadata.get("source", "Unknown file")
        location = display_page(doc.metadata)
        formatted.append(f"[Source: {source} | Page/location: {location}]\n{doc.page_content}")
    return "\n\n".join(formatted)


def conversation_history(messages: list[dict[str, Any]], limit: int = 10) -> list[tuple[str, str]]:
    """Return recent user/assistant turns in the format expected by LangChain."""
    if limit <= 0:
        return []
    return [
        (message["role"], message["content"])
        for message in messages[-limit:]
        if message.get("role") in ("user", "assistant") and "content" in message
    ]

